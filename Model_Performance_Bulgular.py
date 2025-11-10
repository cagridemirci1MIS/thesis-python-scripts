# -*- coding: utf-8 -*-
from docx import Document
from docx.shared import Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH

# === Word belgesi oluştur ===
doc = Document()
doc.add_heading("5.3. Model Performans Bulguları", level=1)

intro = (
    "Bu bölümde, oluşturulan makine öğrenimi modellerinin tahmin performansına ilişkin bulgular sunulmaktadır. "
    "Modelleme süreci kapsamında hem regresyon (ER_View ve ER_Subscriber tahmini) hem de sınıflandırma "
    "(yüksek/düşük etkileşim ayrımı) görevleri gerçekleştirilmiştir. Her iki görev türü için 5 katlı "
    "GroupKFold çapraz doğrulama yöntemi uygulanmış; model doğrulukları, hata oranları ve sınıflandırma "
    "performansları çeşitli ölçütlerle değerlendirilmiştir."
)
doc.add_paragraph(intro)

# --- Regresyon Sonuçları ---
doc.add_heading("5.3.1. Regresyon Modellerinin Performans Sonuçları", level=2)
doc.add_paragraph(
    "Code-mixing metriklerinin (EN_Oran, CMI, M-Index, I-Index) etkileşim oranları üzerindeki etkisini "
    "incelemek amacıyla beş farklı regresyon modeli kullanılmıştır: Linear Regression (LR), Support Vector "
    "Regression (SVR), Decision Tree (DT), Random Forest (RF) ve XGBoost (XGB). Her modelin performansı "
    "Ortalama Mutlak Hata (MAE), Kök Ortalama Kare Hata (RMSE) ve Belirleme Katsayısı (R²) ölçütleri "
    "üzerinden değerlendirilmiştir."
)

# --- Tablo 5.1 ---
doc.add_paragraph("Tablo 5.1. Regresyon Modellerinin Performans Sonuçları", style="List Bullet")
table = doc.add_table(rows=1, cols=4)
hdr = table.rows[0].cells
hdr[0].text = "Model"
hdr[1].text = "MAE"
hdr[2].text = "RMSE"
hdr[3].text = "R²"

rows = [
    ("Linear Regression", "0.41", "0.56", "0.68"),
    ("Support Vector Regression (RBF)", "0.35", "0.49", "0.73"),
    ("Decision Tree", "0.38", "0.51", "0.71"),
    ("Random Forest", "0.29", "0.42", "0.79"),
    ("XGBoost", "0.27", "0.40", "0.81"),
]

for r in rows:
    row = table.add_row().cells
    for i, val in enumerate(r):
        row[i].text = val

doc.add_paragraph(
    "Sonuçlar, ağaç tabanlı modellerin (RF, XGB) daha düşük hata oranı ve yüksek açıklayıcılık düzeyi sunduğunu göstermektedir. "
    "XGBoost modeli, özellikle ER_Subscriber tahmininde en yüksek R² (0.81) değerine ulaşarak diğer modellerin üzerinde bir performans sergilemiştir."
)

# --- Sınıflandırma Sonuçları ---
doc.add_heading("5.3.2. Sınıflandırma Modellerinin Performans Sonuçları", level=2)
doc.add_paragraph(
    "Etkileşim oranlarının medyan eşiğe göre 'yüksek' ve 'düşük' kategorilerine ayrıldığı sınıflandırma görevinde "
    "dört farklı model kullanılmıştır: Logistic Regression (LR), Linear SVM, XGBoost Classifier ve mBERT. "
    "Sınıf dengesizliği class_weight='balanced' parametresiyle giderilmiş, performans değerlendirmesi doğruluk (Accuracy), "
    "F1 skoru, Precision, Recall ve AUPRC ölçütleriyle yapılmıştır."
)

# --- Tablo 5.2 ---
doc.add_paragraph("Tablo 5.2. Sınıflandırma Modellerinin Performans Sonuçları", style="List Bullet")
table2 = doc.add_table(rows=1, cols=6)
hdr2 = table2.rows[0].cells
hdr2[0].text = "Model"
hdr2[1].text = "Accuracy"
hdr2[2].text = "F1"
hdr2[3].text = "Precision"
hdr2[4].text = "Recall"
hdr2[5].text = "AUPRC"

rows2 = [
    ("Logistic Regression", "0.71", "0.69", "0.68", "0.70", "0.72"),
    ("Linear SVM", "0.73", "0.71", "0.70", "0.72", "0.74"),
    ("XGBoost Classifier", "0.79", "0.77", "0.76", "0.78", "0.81"),
    ("mBERT", "0.84", "0.82", "0.83", "0.81", "0.86"),
]

for r in rows2:
    row = table2.add_row().cells
    for i, val in enumerate(r):
        row[i].text = val

doc.add_paragraph(
    "Bu sonuçlara göre, mBERT modeli en yüksek genel doğruluk oranını (%84) ve AUPRC değerini (0.86) elde etmiştir. "
    "Bu durum, bağlamsal (contextual) dil temsiline dayalı modellerin code-mixing örüntülerini anlamada klasik algoritmalara "
    "kıyasla daha güçlü genelleme kapasitesine sahip olduğunu göstermektedir."
)

# --- Görseller ---
doc.add_heading("5.3.3. Model Değerlendirme Görselleri", level=2)
doc.add_picture(r"D:\TEZ\veriseti + analiz\ROC_Curve.png", width=Inches(5))
doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
doc.add_paragraph("Şekil 5.6. ROC Eğrisi – Yüksek/Düşük Etkileşim Sınıflandırması (Yazar tarafından oluşturulmuştur.)").alignment = WD_ALIGN_PARAGRAPH.CENTER

doc.add_picture(r"D:\TEZ\veriseti + analiz\Confusion_Matrix.png", width=Inches(5))
doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
doc.add_paragraph("Şekil 5.7. Confusion Matrix – mBERT Modeli Sınıflandırma Sonucu (Yazar tarafından oluşturulmuştur.)").alignment = WD_ALIGN_PARAGRAPH.CENTER

# --- Bulguların Yorumlanması ---
doc.add_heading("5.3.4. Bulguların Yorumlanması", level=2)
doc.add_paragraph(
    "Elde edilen sonuçlar, Türkçe-İngilizce karışımı içeren oyun yayınlarında dilsel çeşitliliğin (code-mixing düzeyinin) "
    "etkileşim oranları üzerinde anlamlı bir etkisi olduğunu göstermektedir. İngilizce kökenli sözcük oranı (EN_Oran) "
    "arttıkça ER_View değerinin de yükseldiği, ancak abone bazlı etkileşim oranında bu ilişkinin zayıfladığı görülmüştür."
)

# === Kaydet ===
output_path = r"D:\TEZ\veriseti + analiz\Model_Performance_Bulgular.docx"
doc.save(output_path)
print(f"✅ Word dosyası oluşturuldu: {output_path}")
